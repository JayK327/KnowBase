# src/ingestion/extractors/docx_extractor.py
"""
DOCX Extractor (python-docx)
============================

Converts DOCX files into clean, structured text while preserving:
- Heading hierarchy (H1/H2/H3 → Markdown # / ## / ###)
- Paragraph flow
- Tables (converted into Markdown format)

Also extracts document metadata like author and creation date from core properties.
"""

import io
import logging
from docx import Document
from src.ingestion.models import RawDocument, ExtractedDocument, DocumentFormat

logger = logging.getLogger(__name__)


class DocxExtractor:
    def _table_to_markdown(self, table) -> str:
        rows = []
        for i, row in enumerate(table.rows):
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")

            # Add markdown separator row after header
            if i == 0:
                rows.append("| " + " | ".join(["---"] * len(cells)) + " |")

        return "\n".join(rows)

    def extract(self, doc: RawDocument) -> ExtractedDocument:
        document = Document(io.BytesIO(doc.raw_bytes))
        parts, title = [], None

        for element in document.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                for para in document.paragraphs:
                    if para._element is element:
                        text = para.text.strip()
                        if not text:
                            continue

                        style = para.style.name if para.style else ""

                        if "Heading 1" in style:
                            if not title:
                                title = text
                            parts.append(f"\n\n# {text}\n")
                        elif "Heading 2" in style:
                            parts.append(f"\n\n## {text}\n")
                        elif "Heading 3" in style:
                            parts.append(f"\n\n### {text}\n")
                        else:
                            parts.append(text)
                        break

            elif tag == "tbl":
                for tbl in document.tables:
                    if tbl._element is element:
                        parts.append("\n\n" + self._table_to_markdown(tbl) + "\n")
                        break

        text = "\n\n".join(p for p in parts if p.strip())
        props = document.core_properties

        return ExtractedDocument(
            doc_id=doc.doc_id,
            source_path=doc.source_path,
            text=text,
            format=DocumentFormat.DOCX,
            title=title,
            author=props.author,
            created_at=props.created.isoformat() if props.created else None,
            metadata=doc.metadata,
        )
